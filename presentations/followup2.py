# -*- coding: utf-8 -*-
"""Follow-up 2 (October 26, 2026), on the department's shared structure.

Starts from V3 so the title, value, scope, plan and closing slides stay as
the supervisor left them, then rebuilds slides 5 to 12 in the order the
standard sets (presentations/estandar/estandar-becarios.html):

   5 Since the last review     9 Project
   6 Gantt                     10 Results
   7 ILUO status               11 Next steps
   8 What I learned            12 Conclusions

    python3 followup2.py V3.pptx template.pptx
"""
import io
import os
import sys
import zipfile

from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC, TEMPLATE = sys.argv[1], sys.argv[2]
OUT = os.path.join(HERE, "Tenneco_FollowUp2_RoboticWelding_Diego_de_Leon.pptx")
GLASS_DIR = os.path.join(HERE, "build", "glass_fu2")

NAVY  = RGBColor(0x05, 0x1C, 0x2C)
LIME  = RGBColor(0xD5, 0xFB, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0xC6, 0xCD, 0xD1)
DIM   = RGBColor(0x8A, 0x9B, 0xA8)
BLUE  = RGBColor(0x00, 0x33, 0xA0)
SLATE = RGBColor(0x6B, 0x7A, 0x83)
REAL  = RGBColor(0x00, 0xB0, 0x50)       # plant schedule colours
PLAN  = RGBColor(0x9D, 0xC3, 0xE6)
# commitment states, the four words the standard allows
STATE = {"Done": REAL, "On track": BLUE,
         "At risk": RGBColor(0xD9, 0x8E, 0x04),
         "Late": RGBColor(0xC6, 0x28, 0x28)}
FONT = "Segoe UI"
SW = 13.333

# ---------------------------------------------------------------- theme
class T:
    light = False
    txt, sub, acc, mut = WHITE, GRAY, LIME, DIM

def theme(light):
    T.light = light
    if light:
        T.txt, T.sub, T.acc, T.mut = NAVY, SLATE, BLUE, SLATE
    else:
        T.txt, T.sub, T.acc, T.mut = WHITE, GRAY, LIME, DIM

# ---------------------------------------------------------------- helpers
def txbox(s, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tb, tf

def para(tf, t, size, color, bold=False, first=False, space_before=0,
         align=None, spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(0)
    if spacing:
        p.line_spacing = spacing
    r = p.add_run()
    r.text = t
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return p

def text(s, x, y, w, h, t, size, color, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    tb, tf = txbox(s, x, y, w, h, align=align, anchor=anchor)
    para(tf, t, size, color, bold=bold, first=True, spacing=spacing)
    return tb

def label(s, x, y, t, w=6.0):
    return text(s, x, y, w, 0.3, t, 16, T.acc, bold=True)

def _alpha(parent, pct):
    parent.find(qn('a:srgbClr')).append(parse_xml(
        '<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/'
        'main" val="%d"/>' % int(pct * 1000)))

def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, opacity=100,
         line_opacity=100, shape=MSO_SHAPE.RECTANGLE, name=None, dash=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
        if opacity < 100:
            _alpha(sp._element.spPr.find(qn('a:solidFill')), opacity)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
        ln = sp._element.spPr.find(qn('a:ln'))
        if line_opacity < 100:
            _alpha(ln.find(qn('a:solidFill')), line_opacity)
        if dash:
            ln.append(parse_xml(
                '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/'
                'drawingml/2006/main" val="%s"/>' % dash))
    st = sp._element.find(qn('p:style'))
    if st is not None:
        sp._element.remove(st)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sp

def chevron(s, x, y):
    rect(s, x, y, 0.14, 0.14, fill=T.acc, shape=MSO_SHAPE.CHEVRON)

def rule(s, x, y, w=1.20):
    rect(s, x, y, w, 0.03, fill=T.acc)

# Glass exactly as in deck.py and revise_v3.py: one rectangle, filled on dark
# slides with the piece of road it covers (blurred, lifted, saturated, 50 %
# opacity), and a soft navy tint on white. Same shape kind on both themes, so
# Morph pairs the panels from slide to slide.
_tz = zipfile.ZipFile(TEMPLATE)
ROAD = Image.open(io.BytesIO(_tz.read("ppt/media/image4.jpg"))).convert("RGB")
PX = 1920 / SW
_gn = [0]

def glass(s, x, y, w, h, name=None):
    if T.light:
        return rect(s, x, y, w, h, fill=NAVY, opacity=4, line=NAVY,
                    line_opacity=12, name=name)
    crop = ROAD.crop((round(x * PX), round(y * PX),
                      round((x + w) * PX), round((y + h) * PX)))
    crop = crop.filter(ImageFilter.GaussianBlur(5))
    crop = Image.blend(crop, Image.new("RGB", crop.size, (98, 152, 192)), 0.5)
    crop = ImageEnhance.Color(crop).enhance(1.25)
    os.makedirs(GLASS_DIR, exist_ok=True)
    _gn[0] += 1
    path = os.path.join(GLASS_DIR, "glass_%02d.jpg" % _gn[0])
    crop.save(path, quality=90)
    sp = rect(s, x, y, w, h, line=WHITE, line_opacity=16, name=name)
    _, rid = s.part.get_or_add_image_part(path)
    spPr = sp._element.spPr
    spPr.remove(spPr.find(qn('a:noFill')))
    spPr.find(qn('a:prstGeom')).addnext(parse_xml(
        '<a:blipFill xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        '2006/main" xmlns:r="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships" rotWithShape="1">'
        '<a:blip r:embed="%s"><a:alphaModFix amt="50000"/></a:blip>'
        '<a:stretch><a:fillRect/></a:stretch></a:blipFill>' % rid))
    return sp

def photo(s, x, y, w, h, caption, name):
    """Space for one of Diego's photos: delete the label, drop the photo on
    top and crop it to the frame."""
    glass(s, x, y, w, h, name=name)
    rect(s, x + 0.15, y + 0.15, w - 0.30, h - 0.30, line=T.sub,
         line_opacity=35)
    tb, tf = txbox(s, x, y + h / 2 - 0.28, w, 0.6, align=PP_ALIGN.CENTER)
    para(tf, "PHOTO", 11, T.acc, bold=True, first=True)
    para(tf, caption, 12, T.sub, space_before=2, align=PP_ALIGN.CENTER)

def notes(s, *points):
    s.notes_slide.notes_text_frame.text = "\n".join("• " + p for p in points)

# ---------------------------------------------------------------- deck
prs = Presentation(SRC)
layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}
RELNS = qn('r:id')
sld_lst = prs.slides._sldIdLst
ids = list(sld_lst)
# V3: keep 1 title, 2 value, 3 scope, 4 plan and 14 closing
for k, sid in enumerate(ids):
    if k in (0, 1, 2, 3, 13):
        continue
    prs.part.drop_rel(sid.get(RELNS))
    sld_lst.remove(sid)

def new_slide(title, sub, light):
    theme(light)
    s = prs.slides.add_slide(
        layouts["Standard_Light" if light else "Title Image_blank"])
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    for idx, t, size, color, bold, nm in (
            (20, title, 28, T.txt, True, "!! Title 1"),
            (21, sub, 13, T.acc, False, "!! Subtitle 1")):
        tf = ph[idx].text_frame
        tf.text = t
        r = tf.paragraphs[0].runs[0]
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = FONT; r.font.color.rgb = color
        ph[idx].name = nm
    return s

S = list(prs.slides)

# ============================================================ 1 · title
ph = {p.placeholder_format.idx: p for p in S[0].placeholders}
p0 = ph[14].text_frame.paragraphs[0]
for r in p0.runs[1:]:
    r._r.getparent().remove(r._r)
p0.runs[0].text = "Follow-up 2  |  October 26, 2026"
notes(S[0],
      "Second review. Same program: Robotics Setter Technician.",
      "Today: what I promised last time, what happened, and what comes next.")

# ============================================================ 5 · since the last review
s = new_slide("The plan slipped where the floor needed me",
              "Since the last review  |  Commitments from September 25", True)
glass(s, 0.28, 1.55, 12.77, 3.95, name="!! Glass A")
for x, w, h in ((0.62, 6.2, "WHAT I COMMITTED TO"), (7.10, 1.6, "STATE"),
                (8.90, 4.0, "WHAT HAPPENED")):
    text(s, x, 1.75, w, 0.25, h, 9, SLATE, bold=True)
ROWS5 = [
    ("Close the four open Level I topics", "Done",
     "Closed. Theory keeps growing in parallel with the other levels."),
    ("Pass the Level I exam in week 40", "Late",
     "Not taken yet. Date to be set."),
    ("Validate Level L on the floor in October", "Late",
     "I already do all five skills on my own. Date to be set."),
    ("Reach Levels U and O before January", "On track",
     "All three U skills and four of six O skills are moving."),
    ("Keep measuring 4M and find the next improvement", "On track",
     "4M ran the DOC and DPF PV parts. The search continues."),
]
for k, (what, st, note) in enumerate(ROWS5):
    y = 2.08 + k * 0.66
    if k:
        rect(s, 0.62, y - 0.12, 12.10, 0.01, fill=NAVY, opacity=10)
    text(s, 0.62, y, 6.2, 0.5, what, 14, NAVY, bold=True)
    pill = rect(s, 7.10, y - 0.01, 1.30, 0.32, fill=STATE[st],
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pill.adjustments[0] = 0.5
    pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(pill.text_frame, st, 11, WHITE, bold=True, first=True,
         align=PP_ALIGN.CENTER)
    text(s, 8.90, y + 0.01, 4.0, 0.55, note, 12, SLATE)
glass(s, 0.28, 5.68, 12.77, 0.80, name="!! Strip")
text(s, 0.62, 5.95, 1.6, 0.25, "WHY", 11, BLUE, bold=True)
text(s, 2.05, 5.84, 10.7, 0.55, "In October the industrialization project "
     "needed the whole team: parts already released failed metrology, and "
     "the SCR model started. The exams waited; the floor did not.", 12, NAVY)
notes(s,
      "Be direct: two commitments are late, and say why.",
      "One done (Level I topics), two on track (U and O, 4M).",
      "Exam and floor validation: dates to be set with my supervisor.")

# ============================================================ 6 · Gantt
s = new_slide("Skills kept moving while the exams waited",
              "Gantt  |  Technical and practical skills, done and next", False)
glass(s, 0.28, 1.50, 12.77, 4.55, name="!! Glass A")
GX, GW, NW = 3.45, 9.40, 23                 # Aug 3, 2026 to Jan 10, 2027
WK = GW / NW
gx = lambda w: GX + w * WK
for m, a, b in (("August", 0, 5), ("September", 5, 9), ("October", 9, 13),
                ("November", 13, 18), ("December", 18, 22), ("Jan", 22, 23)):
    text(s, gx(a), 1.60, (b - a) * WK, 0.25, m, 11, WHITE, bold=True,
         align=PP_ALIGN.CENTER)
    if a:
        rect(s, gx(a), 1.62, 0.01, 4.10, fill=WHITE, opacity=18)
for i in range(NW):
    text(s, gx(i), 1.86, WK, 0.2, str(32 + i if i < 22 else 1), 8, DIM,
         align=PP_ALIGN.CENTER)
rect(s, 0.50, 2.10, 12.35, 0.01, fill=WHITE, opacity=25)

TODAY = 12.1                                # Monday, October 26 (week 44)
G_ROWS = [
    ("Level I, theory", "In parallel with every level", (0, TODAY), (TODAY, 23), None),
    ("Level I exam", "Moved from week 40, date to be set", None, None, "moved"),
    ("Level L, on my own", "Floor validation, date to be set", (0, TODAY), None, None),
    ("Level U, programming", "SKS programs and parameters", (5, TODAY), (TODAY, 23), None),
    ("Level O, improve and teach", "Training team members, problem solving", (0, TODAY), (TODAY, 23), None),
    ("DOC and DPF, 95%", "PVs in 4M, PPAP in week 1", (0, TODAY), (TODAY, 22), 22.5),
    ("Metrology recovery", "DPF parallelism, 8 parts remade", (8, 11), None, None),
    ("SCR, new model", "PVs of two models, released", (8, TODAY), (TODAY, 23), None),
]
for k, (lab, det, real, plan, mile) in enumerate(G_ROWS):
    y = 2.18 + k * 0.44
    if k:
        rect(s, 0.50, y - 0.04, 12.35, 0.01, fill=WHITE, opacity=8)
    text(s, 0.50, y, 2.9, 0.22, lab, 11, WHITE, bold=True)
    text(s, 0.50, y + 0.20, 2.9, 0.2, det, 8.5, GRAY)
    if plan:
        rect(s, gx(plan[0]), y + 0.08, (plan[1] - plan[0]) * WK, 0.22, fill=PLAN)
    if real:
        rect(s, gx(real[0]), y + 0.08, (real[1] - real[0]) * WK, 0.22, fill=REAL)
    if mile == "moved":
        rect(s, gx(8.5) - 0.11, y + 0.08, 0.22, 0.22, line=GRAY, lw=1.25,
             shape=MSO_SHAPE.DIAMOND)
        text(s, gx(8.5) + 0.18, y + 0.08, 2.2, 0.22, "Week 40, moved", 9, GRAY)
    elif mile is not None:
        rect(s, gx(mile) - 0.11, y + 0.08, 0.22, 0.22, fill=WHITE,
             shape=MSO_SHAPE.DIAMOND)
        text(s, gx(mile) - 1.05, y + 0.08, 0.85, 0.22, "PPAP", 9, NAVY,
             bold=True, align=PP_ALIGN.RIGHT)
rect(s, gx(TODAY), 2.12, 0.0, 3.55, line=LIME, lw=1.5, dash="dash")
text(s, gx(TODAY) - 0.7, 5.72, 1.4, 0.22, "Today, Oct 26", 9, LIME, bold=True,
     align=PP_ALIGN.CENTER)
lx = 3.45
for t, kw in (("Real", dict(fill=REAL)), ("Plan", dict(fill=PLAN)),
              ("Milestone", dict(fill=WHITE, shape=MSO_SHAPE.DIAMOND)),
              ("Moved", dict(line=GRAY, lw=1.25, shape=MSO_SHAPE.DIAMOND))):
    rect(s, lx, 5.77, 0.16, 0.16, **kw)
    text(s, lx + 0.24, 5.74, 1.0, 0.22, t, 9, GRAY)
    lx += 1.05
notes(s,
      "Green is real, light blue is plan, same as the plant Excel.",
      "Weeks 40 to 42: metrology recovery and SCR start, in parallel.",
      "Level I theory runs in parallel; the exam moved from week 40.",
      "DOC and DPF at 95%, PPAP in the first week of January.")

# ============================================================ 7 · ILUO status
s = new_slide("15 of 23 skills achieved, only 2 not started",
              "ILUO status  |  September 25: 11 achieved, 6 not started", True)
glass(s, 0.28, 1.75, 8.25, 4.25, name="!! Glass A")
ROWS7 = [("I", "Understands", 9, 0, 0), ("L", "With support", 5, 0, 0),
         ("U", "Alone", 0, 3, 0), ("O", "Improves", 1, 3, 2)]
for k, (lv, nm, own, wip, todo) in enumerate(ROWS7):
    y = 2.12 + k * 0.84
    text(s, 0.68, y - 0.06, 0.5, 0.5, lv, 24, BLUE, bold=True)
    text(s, 1.22, y - 0.01, 2.2, 0.3, nm, 13, NAVY, bold=True)
    text(s, 1.22, y + 0.26, 2.2, 0.3, "%d skills" % (own + wip + todo), 11, SLATE)
    for j in range(own + wip + todo):
        x = 3.45 + j * 0.40
        if j < own:
            rect(s, x, y + 0.06, 0.26, 0.26, fill=BLUE, shape=MSO_SHAPE.OVAL)
        elif j < own + wip:
            rect(s, x, y + 0.06, 0.26, 0.26, fill=BLUE, opacity=35, line=BLUE,
                 lw=1.25, shape=MSO_SHAPE.OVAL)
        else:
            rect(s, x, y + 0.06, 0.26, 0.26, line=SLATE, lw=1.0,
                 line_opacity=60, shape=MSO_SHAPE.OVAL)
lx = 3.45
for t, kw in (("Achieved", dict(fill=BLUE)),
              ("In progress", dict(fill=BLUE, opacity=35, line=BLUE, lw=1.25)),
              ("Not started", dict(line=SLATE, lw=1.0, line_opacity=60))):
    rect(s, lx, 5.52, 0.18, 0.18, shape=MSO_SHAPE.OVAL, **kw)
    text(s, lx + 0.28, 5.49, 1.4, 0.25, t, 11, SLATE)
    lx += 1.62
glass(s, 8.85, 1.75, 4.20, 4.25, name="!! Glass B")
label(s, 9.25, 2.05, "WHAT MOVED", w=3.6)
MOVED = [("I", "Four topics closed; theory keeps growing"),
         ("L", "All five skills, now on my own"),
         ("U", "SKS programs and parameters started"),
         ("O", "Training team members, problem-solving basics")]
for k, (lv, t) in enumerate(MOVED):
    y = 2.55 + k * 0.66
    text(s, 9.25, y - 0.02, 0.4, 0.3, lv, 14, BLUE, bold=True)
    text(s, 9.62, y, 3.25, 0.6, t, 13, NAVY)
text(s, 9.25, 5.35, 3.6, 0.55, "Still pending: the Level I exam and the "
     "Level L floor validation.", 11, BLUE, bold=True)
notes(s,
      "From 11 to 15 achieved, from 6 to 2 not started.",
      "Not started: GD&T and Lean.",
      "Level I and L achieved, but the formal exam and validation are pending.")

# ============================================================ 8 · what I learned
s = new_slide("October’s firefight taught me more than the plan",
              "What I learned  |  Weeks 40 to 42", False)
glass(s, 0.28, 1.75, 6.10, 3.55, name="!! Glass A")
label(s, 0.68, 2.05, "WHAT HAPPENED")
tb, tf = txbox(s, 0.68, 2.50, 5.35, 2.7)
para(tf, "DPF parts that were already released started failing "
     "metrology. A metrology expert measured everything again: the main "
     "failure was parallelism.", 14, WHITE, first=True, spacing=1.1)
para(tf, "At the same time the SCR model started: bigger, heavier and "
     "harder to build, with more machines and processes.", 14, GRAY,
     space_before=12, spacing=1.1)
glass(s, 6.70, 1.75, 6.35, 3.55, name="!! Glass B")
label(s, 7.10, 2.05, "WHAT I LEARNED ON THE FLOOR")
LEARN = ["Dimensional adjustments on welding fixtures",
         "Tuning welding parameters until the part is right",
         "Reviewing the process step by step",
         "Training team members on machines, adjustments, process and product",
         "The basics of problem-solving tools"]
for k, t in enumerate(LEARN):
    y = 2.55 + k * 0.52
    chevron(s, 7.10, y + 0.07)
    text(s, 7.40, y, 5.40, 0.5, t, 13, WHITE)
glass(s, 0.28, 5.42, 12.77, 0.62, name="!! Strip")
text(s, 0.62, 5.62, 2.2, 0.25, "HOW", 11, LIME, bold=True)
text(s, 2.75, 5.60, 10.2, 0.3, "On the floor, alongside Engineering and the "
     "metrology expert, while the line kept running.", 12, WHITE)
notes(s,
      "Parallelism: two faces of the part that must stay parallel, within "
      "tolerance.",
      "The firefight was where U and O skills grew: adjustments, parameters, "
      "training people.",
      "SCR: bigger, heavier, more processes than DOC and DPF.")

# ============================================================ 9 · project
s = new_slide("Industrialization: DOC and DPF at 95%, SCR started",
              "Project  |  MY27 Cummins program, together with Engineering", True)
glass(s, 0.28, 1.75, 6.35, 4.25, name="!! Glass A")
BLOCKS = [("THE PROBLEM", "DPF parts already released failed metrology: "
           "parallelism out of tolerance."),
          ("WHAT WE DID", "A metrology expert measured everything again, and "
           "we remade 8 DPF parts: 4 short and 4 long canning."),
          ("WHERE IT STANDS", "DOC and DPF at 95%, only documentation left. "
           "SCR started in October: 7 PV parts of each of two models, "
           "released.")]
for k, (lab, d) in enumerate(BLOCKS):
    y = 2.05 + k * 1.30
    label(s, 0.68, y, lab)
    tb, tf = txbox(s, 0.68, y + 0.40, 5.6, 0.85)
    para(tf, d, 13, NAVY, first=True, spacing=1.1)
photo(s, 6.98, 1.75, 6.07, 2.02, "DPF parts or the metrology check",
      "!! Glass B")
photo(s, 6.98, 3.98, 6.07, 2.02, "The SCR model", "!! Glass C")
notes(s,
      "Problem: parallelism on DPF parts that were already released.",
      "Action: full re-measurement with a metrology expert, 8 DPF remade.",
      "Status: DOC and DPF 95%, only documents; SCR started, 14 PVs released.")

# ============================================================ 10 · results
s = new_slide("DOC and DPF moved from 90% to 95%",
              "Results  |  September 25 to October 26", False)
glass(s, 0.28, 1.75, 8.25, 4.25, name="!! Glass A")
label(s, 0.68, 2.05, "DOC AND DPF, READY TO LAUNCH")
for k, (when, pct, hot) in enumerate((("September 25", 90, False),
                                      ("October 26", 95, True))):
    y = 2.55 + k * 0.62
    text(s, 0.68, y + 0.02, 1.7, 0.3, when, 12, WHITE if hot else GRAY,
         bold=hot)
    rect(s, 2.45, y + 0.04, 4.80, 0.26, fill=GRAY, opacity=25)
    rect(s, 2.45, y + 0.04, 4.80 * pct / 100, 0.26,
         fill=LIME if hot else GRAY, opacity=100 if hot else 70)
    text(s, 7.38, y, 1.0, 0.35, "%d%%" % pct, 16, LIME if hot else WHITE,
         bold=True)
text(s, 0.68, 3.82, 7.4, 0.3, "Only documentation left. PPAP in the first "
     "week of January.", 12, GRAY)
KPI10 = [("17", "PV parts in 4M", "9 DOC, 4 DPF short, 4 DPF long"),
         ("8", "DPF parts remade", "after re-measuring parallelism"),
         ("14", "SCR PV parts", "two models, 7 each, released")]
for i, (big, lab, det) in enumerate(KPI10):
    x = 0.68 + i * 2.55
    text(s, x, 4.30, 2.4, 0.6, big, 30, LIME, bold=True)
    text(s, x, 4.92, 2.4, 0.3, lab, 12, WHITE, bold=True)
    text(s, x, 5.20, 2.4, 0.5, det, 10, GRAY)
glass(s, 8.85, 1.75, 4.20, 4.25, name="!! Glass B")
label(s, 9.25, 2.05, "4M, STILL RUNNING", w=3.6)
text(s, 9.25, 2.55, 3.6, 0.8, "93 s", 40, LIME, bold=True)
text(s, 9.25, 3.35, 3.6, 0.3, "per part in the robotic cell", 12, WHITE)
text(s, 9.25, 3.95, 3.6, 0.6, "−47%", 28, LIME, bold=True)
text(s, 9.25, 4.55, 3.6, 0.5, "cycle time, from 174 s welding by hand", 12,
     GRAY)
text(s, 9.25, 5.35, 3.6, 0.4, "Data: engineering team, October 2026", 9, DIM)
notes(s,
      "Before and after: 90% on September 25, 95% today.",
      "17 PVs ran in the 4M cell, the improvement from last time.",
      "8 DPF remade after the metrology check; 14 SCR PVs already released.")

# ============================================================ 11 · next steps
s = new_slide("Next: set the exam dates, then Level U",
              "Next steps  |  Until January", True)
glass(s, 0.28, 1.75, 7.90, 4.25, name="!! Glass A")
label(s, 0.68, 2.05, "WHAT COMES NEXT")
NEXT = [("Level I exam and Level L floor validation", "Dates to be set with my supervisor"),
        ("Level U: SKS programs and welding parameters", "November to December"),
        ("Level O: problem solving and the next improvement", "November to December"),
        ("DOC and DPF PPAP", "First week of January")]
for k, (t, when) in enumerate(NEXT):
    y = 2.55 + k * 0.82
    if k:
        rect(s, 0.68, y - 0.14, 7.10, 0.01, fill=NAVY, opacity=10)
    text(s, 0.68, y, 7.1, 0.3, t, 14, NAVY, bold=True)
    text(s, 0.68, y + 0.33, 7.1, 0.3, when, 12, SLATE)
glass(s, 8.55, 1.75, 4.50, 4.25, name="!! Glass B")
label(s, 8.95, 2.05, "WHAT I NEED", w=3.8)
NEED = ["Dates for my Level I exam and Level L validation",
        "Supervised time programming in SKS and Miller",
        "Feedback on the training proposal"]
for k, t in enumerate(NEED):
    y = 2.60 + k * 0.85
    chevron(s, 8.95, y + 0.08)
    text(s, 9.25, y, 3.55, 0.75, t, 14, NAVY)
notes(s,
      "Ask clearly for the two dates: they are the next review's commitments.",
      "U is the focus for November and December.",
      "Training proposal: still no comments, ask again.")

# ============================================================ 12 · conclusions
s = new_slide("Conclusions and next steps", "What is in place, and what comes next", False)
glass(s, 0.28, 1.75, 6.35, 4.05, name="!! Glass A")
label(s, 0.68, 2.05, "IN PLACE")
glass(s, 6.98, 1.75, 6.07, 4.05, name="!! Glass B")
label(s, 7.38, 2.05, "NEXT")
INPLACE = ["15 of 23 skills achieved; Levels I and L done on my own.",
           "DOC and DPF at 95%, and SCR started with 14 PV parts released.",
           "A month of real problem solving: fixtures, parameters and training."]
NEXTS = ["Set and pass the Level I exam and the Level L validation.",
         "Reach Level U before January.",
         "Support the DOC and DPF PPAP in the first week of January."]
for x, items, w in ((0.68, INPLACE, 5.55), (7.38, NEXTS, 5.27)):
    for k, t in enumerate(items):
        text(s, x, 2.55 + k * 1.00, w, 0.8, t, 14, WHITE)
text(s, 3.40, 6.25, 9.6, 0.4, "The plan slipped a few weeks; my skills did "
     "not.", 15, LIME, bold=True)
notes(s,
      "Close with the last line. Pause, then questions.")

# closing slide back to the end
sld_lst.append(sld_lst[4])

prs.save(OUT)
print("saved", OUT)
