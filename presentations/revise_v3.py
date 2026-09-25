# -*- coding: utf-8 -*-
"""V2 -> V3: the supervisor's review, applied on top of his own file.

V2 is Diego's deck after his supervisor edited it by hand and Diego placed
his photos, so this script never regenerates the deck. It opens V2 and
changes only what the review asked for:

  1. Slide 5: the timeline becomes a Gantt, as the corporate template asks
     ("Explicar en un Gantt las habilidades/conocimientos técnicos y
     prácticos hasta el momento y siguientes pasos").
  2. Level I: which topics are mastered and which are still open.
  3. Slide 4: Diego's ILUO matrix, paraphrased as the plan to reach the
     program objective.

It also lands the supervisor's wording (plain punctuation, no long dashes)
and makes the numbers agree from slide to slide.

    python3 revise_v3.py V2.pptx template.pptx
"""
import io
import os
import sys
import zipfile

from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC, TEMPLATE = sys.argv[1], sys.argv[2]
OUT = os.path.join(HERE, "Tenneco_Robotics_Setter_Diego_de_Leon_V3.pptx")
GLASS_DIR = os.path.join(HERE, "build", "glass_v3")

NAVY  = RGBColor(0x05, 0x1C, 0x2C)
LIME  = RGBColor(0xD5, 0xFB, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0xC6, 0xCD, 0xD1)
DIM   = RGBColor(0x8A, 0x9B, 0xA8)
BLUE  = RGBColor(0x00, 0x33, 0xA0)
SLATE = RGBColor(0x6B, 0x7A, 0x83)
# the plant schedule's own colours, so the Gantt reads like the Excel one
REAL  = RGBColor(0x00, 0xB0, 0x50)
PLAN  = RGBColor(0x9D, 0xC3, 0xE6)
FONT  = "Segoe UI"
SW = 13.333

# ---------------------------------------------------------------- helpers
def txbox(s, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tb, tf

def para(tf, t, size, color, bold=False, first=False, space_before=0,
         align=None, italic=False, spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(0)
    if spacing:
        p.line_spacing = spacing
    r = p.add_run()
    r.text = t
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    r.font.color.rgb = color
    return p

def text(s, x, y, w, h, t, size, color, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    tb, tf = txbox(s, x, y, w, h, align=align, anchor=anchor)
    para(tf, t, size, color, bold=bold, first=True, spacing=spacing)
    return tb

def _alpha(parent, pct):
    parent.find(qn('a:srgbClr')).append(parse_xml(
        '<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/'
        'main" val="%d"/>' % int(pct * 1000)))

def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, opacity=100,
         line_opacity=100, shape=MSO_SHAPE.RECTANGLE, name=None, dash=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
        if opacity < 100:
            _alpha(sp._element.spPr.find(qn('a:solidFill')), opacity)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
        ln = sp._element.spPr.find(qn('a:ln'))
        if line_opacity < 100:
            _alpha(ln.find(qn('a:solidFill')), line_opacity)
        if dash:
            ln.append(parse_xml(
                '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/'
                'drawingml/2006/main" val="%s"/>' % dash))
    st = sp._element.find(qn('p:style'))
    if st is not None:
        sp._element.remove(st)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sp

# Glass exactly as in deck.py: one rectangle, filled on dark slides with the
# piece of background it covers (blurred, lifted, saturated, 50 % opacity),
# and a soft navy tint on white. One shape kind on both, so Morph pairs it.
_tz = zipfile.ZipFile(TEMPLATE)
ROAD = Image.open(io.BytesIO(_tz.read("ppt/media/image4.jpg"))).convert("RGB")
PX = 1920 / SW
_gn = [0]

def glass(s, x, y, w, h, light, name=None):
    if light:
        return rect(s, x, y, w, h, fill=NAVY, opacity=4, line=NAVY,
                    line_opacity=12, name=name)
    crop = ROAD.crop((round(x * PX), round(y * PX),
                      round((x + w) * PX), round((y + h) * PX)))
    crop = crop.filter(ImageFilter.GaussianBlur(5))
    crop = Image.blend(crop, Image.new("RGB", crop.size, (98, 152, 192)), 0.5)
    crop = ImageEnhance.Color(crop).enhance(1.25)
    os.makedirs(GLASS_DIR, exist_ok=True)
    _gn[0] += 1
    path = os.path.join(GLASS_DIR, "glass_%02d.jpg" % _gn[0])
    crop.save(path, quality=90)
    sp = rect(s, x, y, w, h, line=WHITE, line_opacity=16, name=name)
    _, rid = s.part.get_or_add_image_part(path)
    spPr = sp._element.spPr
    spPr.remove(spPr.find(qn('a:noFill')))
    spPr.find(qn('a:prstGeom')).addnext(parse_xml(
        '<a:blipFill xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        '2006/main" xmlns:r="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships" rotWithShape="1">'
        '<a:blip r:embed="%s"><a:alphaModFix amt="50000"/></a:blip>'
        '<a:stretch><a:fillRect/></a:stretch></a:blipFill>' % rid))
    return sp

def clear(s):
    """Drop every shape except the title placeholders, and the image
    relationships only those shapes used."""
    keep = ("!! Title 1", "!! Subtitle 1")
    for sh in list(s.shapes):
        if sh.name in keep:
            continue
        for blip in sh._element.iter(qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            if rid:
                s.part.drop_rel(rid)
        sh._element.getparent().remove(sh._element)

def shape(s, name):
    return next(sh for sh in s.shapes if sh.name == name)

def set_title(s, title, sub=None):
    for nm, t in (("!! Title 1", title), ("!! Subtitle 1", sub)):
        if t is None:
            continue
        p0 = shape(s, nm).text_frame.paragraphs[0]
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
        p0.runs[0].text = t

def retext(s, name, *paras):
    """Replace a text box's paragraphs, keeping the first run's look of each
    paragraph (size, colour, weight), so the supervisor's styling survives."""
    tf = shape(s, name).text_frame
    ps = list(tf.paragraphs)
    for i, t in enumerate(paras):
        p = ps[min(i, len(ps) - 1)]
        if i >= len(ps):                       # clone the last paragraph
            new = parse_xml(p._p.xml)
            tf._txBody.append(new)
            p = tf.paragraphs[-1]
        runs = p.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = t
    for p in list(tf.paragraphs)[len(paras):]:
        p._p.getparent().remove(p._p)

def notes(s, *points):
    s.notes_slide.notes_text_frame.text = "\n".join("• " + p for p in points)

prs = Presentation(SRC)
S = list(prs.slides)

# ============================================================ 2 · Win
# Supervisor's sentence, kept in substance, with the stray comma and the
# long dashes gone.
retext(S[1], "TextBox 5",
       "We must earn the trust of our employees and customers.",
       "For me, winning means caring for the ergonomics of our welders while "
       "protecting product quality. As an engineer, I do it by optimizing the "
       "process with robotic systems.")
notes(S[1],
      "Win: earn the trust of our people and customers.",
      "Two things at once: the welder’s ergonomics and the product’s quality.",
      "My way as an engineer: optimize the process with robots (4M).")

# ============================================================ 4 · ILUO plan
s = S[3]
clear(s)
set_title(s, "The plan: four levels, each one validated",
          "My ILUO matrix, paraphrased  |  40 hours of formal training")
tb, tf = txbox(s, 0.28, 1.55, 12.7, 0.3)
para(tf, "OBJECTIVE   ", 11, LIME, bold=True, first=True)
r = tf.paragraphs[0].add_run()
r.text = "Become a specialist in robotic welding and advanced technical support."
r.font.size = Pt(13); r.font.name = FONT; r.font.color.rgb = WHITE

COLS = [(0.50, 1.95, "LEVEL"), (2.55, 5.05, "WHAT IT COVERS"),
        (7.80, 2.05, "HOW IT IS VALIDATED"), (10.05, 0.75, "TIME"),
        (10.85, 2.10, "METHOD")]
for x, w, h in COLS:
    text(s, x, 2.02, w, 0.25, h, 9, DIM, bold=True)

PLAN_ROWS = [
    ("I", "Understands",
     "Safety, GMAW equipment, metal transfer, process variables, "
     "discontinuities, joints, positions, symbols and control documents.",
     "Theory exam, 8.0 to pass", "4 h", "Classroom session"),
    ("L", "With support",
     "Basic robotic and dimensional adjustments, checking settings against "
     "the WPS, finding defects and changing consumables.",
     "Validation on the floor", "8 h", "Robotic cell, on the floor (Gemba)"),
    ("U", "Alone",
     "Creating new welding programs in SKS and setting their parameters "
     "(SKS and Miller).",
     "Practical exam", "8 h", "Robotic cell, on the floor (Gemba)"),
    ("O", "Improves and teaches",
     "Training level 2 and 3 staff, improvement ideas, problem solving "
     "(8D, Ishikawa, 5W+2H), GD&T, ANDON and Lean.",
     "Implementation and training, validated", "20 h",
     "80-hour follow-up: TM, WAVE or PSIF"),
]
NAMES = ("!! Glass A", "!! Card 2", "!! Card 3", "!! Glass B")
for k, (lv, nm, what, how, hrs, method) in enumerate(PLAN_ROWS):
    y = 2.32 + k * 0.90
    glass(s, 0.28, y, 12.77, 0.80, light=False, name=NAMES[k])
    text(s, 0.50, y + 0.13, 0.45, 0.5, lv, 24, LIME, bold=True)
    text(s, 0.98, y + 0.12, 1.5, 0.6, nm, 12, WHITE, bold=True)
    text(s, 2.55, y + 0.12, 5.05, 0.6, what, 11, GRAY, spacing=1.05)
    text(s, 7.80, y + 0.12, 2.00, 0.6, how, 11, WHITE)
    text(s, 10.05, y + 0.10, 0.75, 0.5, hrs, 16, LIME, bold=True)
    text(s, 10.85, y + 0.12, 2.10, 0.6, method, 11, GRAY)
notes(s,
      "This is my ILUO matrix, turned into the plan to reach the objective.",
      "Each level has its own test: exam, floor validation, practical exam, "
      "implementation.",
      "40 hours of formal training in total; the rest happens at the cells.",
      "Supervisor’s point: each level is validated before the next one.")

# ============================================================ 5 · Gantt
s = S[4]
clear(s)
set_title(s, "Level I exam next, in week 40",
          "Gantt  |  Technical and practical skills, done and next")
glass(s, 0.28, 1.55, 12.77, 5.05, light=True, name="!! Glass A")

GX, GW, NW = 3.60, 9.25, 22            # chart origin, width, weeks
WK = GW / NW
def gx(week):                           # week index (0 = Aug 3) -> inches
    return GX + week * WK

MONTHS = [("August", 0, 5), ("September", 5, 9), ("October", 9, 13),
          ("November", 13, 18), ("December", 18, 22)]
for m, a, b in MONTHS:
    text(s, gx(a), 1.72, (b - a) * WK, 0.25, m, 11, NAVY, bold=True,
         align=PP_ALIGN.CENTER)
    if a:
        rect(s, gx(a), 1.72, 0.01, 4.35, fill=NAVY, opacity=15)
for i in range(NW):
    text(s, gx(i), 2.00, WK, 0.2, str(32 + i), 8, SLATE,
         align=PP_ALIGN.CENTER)
rect(s, 0.50, 2.27, 12.35, 0.01, fill=NAVY, opacity=25)

TODAY = 7 + 4 / 7                       # Friday, September 25 (week 39)
G_ROWS = [
    # label, detail, real (from, to), plan (from, to), milestone week
    ("Level I, mastered", "Safety, GMAW, variables, defects, documents",
     (0, TODAY), None, None),
    ("Level I, in progress", "Metal transfer, joints, positions, symbols",
     (0, TODAY), (TODAY, 9), None),
    ("Level I exam", "Theory exam, 8.0 to pass", None, None, 8.5),
    ("Level L, floor skills", "Settings, defects, consumables",
     (0, TODAY), (TODAY, 13), None),
    ("Level U, programming", "New programs, SKS and Miller",
     (5, TODAY), (13, 22), None),
    ("Level O, improve and teach", "ANDON calls and the 4M improvement",
     (0, TODAY), (13, 22), None),
    ("Industrialization", "MY27 Cummins, DOC and DPF program",
     (0, TODAY), (TODAY, 22), None),
]
for k, (lab, det, real, plan, mile) in enumerate(G_ROWS):
    y = 2.42 + k * 0.54
    text(s, 0.50, y, 3.0, 0.25, lab, 12, NAVY, bold=True)
    text(s, 0.50, y + 0.24, 3.0, 0.22, det, 9, SLATE)
    if k:
        rect(s, 0.50, y - 0.06, 12.35, 0.01, fill=NAVY, opacity=8)
    if plan:
        rect(s, gx(plan[0]), y + 0.10, (plan[1] - plan[0]) * WK, 0.24,
             fill=PLAN)
    if real:
        rect(s, gx(real[0]), y + 0.10, (real[1] - real[0]) * WK, 0.24,
             fill=REAL)
    if mile is not None:
        rect(s, gx(mile) - 0.13, y + 0.09, 0.26, 0.26, fill=NAVY,
             shape=MSO_SHAPE.DIAMOND)
        text(s, gx(mile) + 0.20, y + 0.10, 1.6, 0.25, "Week 40", 10, NAVY,
             bold=True)

rect(s, gx(TODAY), 2.30, 0.0, 3.80, line=BLUE, lw=1.5, dash="dash")
text(s, gx(TODAY) - 0.6, 6.14, 1.2, 0.22, "Today, Sep 25", 9, BLUE,
     bold=True, align=PP_ALIGN.CENTER)

lx = 0.50
for lab, kw in (("Real", dict(fill=REAL)), ("Plan", dict(fill=PLAN)),
                ("Milestone", dict(fill=NAVY, shape=MSO_SHAPE.DIAMOND))):
    rect(s, lx, 6.20, 0.18, 0.18, **kw)
    text(s, lx + 0.26, 6.17, 1.0, 0.22, lab, 10, SLATE)
    lx += 1.05
notes(s,
      "Green is what really happened, blue is the plan, same as the Excel.",
      "Level I: five topics mastered, four still open before the exam.",
      "Level I exam in week 40. Level L validated on the floor in October.",
      "U and O in November and December; O already started with 4M.",
      "Industrialization runs the whole period.")

# ============================================================ 6 · Level I
s = S[5]
clear(s)
set_title(s, "Level I: five topics mastered, four still open",
          "What I learned")
glass(s, 0.28, 1.75, 7.90, 3.55, light=False, name="!! Glass A")
text(s, 0.68, 2.02, 6.0, 0.3, "MASTERED", 16, LIME, bold=True)
MASTERED = [
    ("Safety", "Working at the cell without risk, always with LOTOTO "
               "(lockout, tagout, tryout)."),
    ("GMAW equipment", "Every part of a GMAW welding system, and how they "
                       "work together."),
    ("Process variables", "Keeping the essential variables under control, "
                          "because they decide weld quality."),
    ("Discontinuities", "Reading the customer’s weld requirements, spotting "
                        "NG welds and finding their cause."),
    ("Control documents", "WPS, PQR and parameter sheets: the approved "
                          "recipe for every weld."),
]
for k, (t, d) in enumerate(MASTERED):
    y = 2.46 + k * 0.55
    if k:
        rect(s, 0.68, y - 0.09, 7.10, 0.01, fill=WHITE, opacity=14)
    text(s, 0.68, y, 2.1, 0.3, t, 13, WHITE, bold=True)
    text(s, 2.85, y + 0.01, 5.05, 0.45, d, 11, GRAY)

glass(s, 8.55, 1.75, 4.50, 3.55, light=False, name="!! Glass B")
text(s, 8.95, 2.02, 3.9, 0.3, "STILL OPEN", 16, LIME, bold=True)
OPEN = ["Metal transfer modes", "Joint geometry", "Welding positions",
        "Welding symbols"]
for k, t in enumerate(OPEN):
    y = 2.52 + k * 0.50
    rect(s, 8.95, y + 0.08, 0.14, 0.14, fill=LIME, shape=MSO_SHAPE.CHEVRON)
    text(s, 9.25, y, 3.6, 0.3, t, 14, WHITE)
text(s, 8.95, 4.62, 3.9, 0.5, "I close them before the exam in week 40.",
     12, GRAY)

glass(s, 0.28, 5.42, 12.77, 0.62, light=False, name="!! Strip")
text(s, 0.62, 5.62, 2.2, 0.25, "HOW I LEARNED IT", 11, LIME, bold=True)
text(s, 2.75, 5.60, 10.2, 0.3, "616 pages of manual and standard, a practice "
     "exam, and daily failures at the cells.", 12, WHITE)
notes(s,
      "Supervisor’s point: be clear that Level I is not complete yet.",
      "Mastered: safety with LOTOTO, GMAW equipment, variables, "
      "discontinuities, control documents.",
      "Still open: transfer modes, joints, positions, symbols.",
      "Plan: close the four before the exam in week 40.")

# ============================================================ 7 · progress
# (level, name, achieved, in progress, not started). Level I follows Diego;
# Level L follows the supervisor, who marked all five skills as achieved.
s = S[6]
clear(s)
set_title(s, "Eleven skills achieved, six in progress",
          "Progress on the 23 skills of the program")
glass(s, 0.28, 1.75, 8.25, 4.25, light=True, name="!! Glass A")
ROWS = [("I", "Understands", 5, 4, 0), ("L", "With support", 5, 0, 0),
        ("U", "Alone", 0, 1, 2), ("O", "Improves", 1, 1, 4)]
for k, (lv, nm, own, wip, todo) in enumerate(ROWS):
    y = 2.12 + k * 0.84
    text(s, 0.68, y - 0.06, 0.5, 0.5, lv, 24, BLUE, bold=True)
    text(s, 1.22, y - 0.01, 2.2, 0.3, nm, 13, NAVY, bold=True)
    text(s, 1.22, y + 0.26, 2.2, 0.3, "%d skills" % (own + wip + todo), 11,
         SLATE)
    for j in range(own + wip + todo):
        x = 3.45 + j * 0.40
        if j < own:
            rect(s, x, y + 0.06, 0.26, 0.26, fill=BLUE, shape=MSO_SHAPE.OVAL)
        elif j < own + wip:
            rect(s, x, y + 0.06, 0.26, 0.26, fill=BLUE, opacity=35, line=BLUE,
                 lw=1.25, shape=MSO_SHAPE.OVAL)
        else:
            rect(s, x, y + 0.06, 0.26, 0.26, line=SLATE, lw=1.0,
                 line_opacity=60, shape=MSO_SHAPE.OVAL)
lx = 3.45
for t, kw in (("Achieved", dict(fill=BLUE)),
              ("In progress", dict(fill=BLUE, opacity=35, line=BLUE,
                                   lw=1.25)),
              ("Not started", dict(line=SLATE, lw=1.0, line_opacity=60))):
    rect(s, lx, 5.52, 0.18, 0.18, shape=MSO_SHAPE.OVAL, **kw)
    text(s, lx + 0.28, 5.49, 1.4, 0.25, t, 11, SLATE)
    lx += 1.62

glass(s, 8.85, 1.75, 4.20, 4.25, light=True, name="!! Glass B")
text(s, 9.25, 2.08, 3.7, 0.3, "ON THE FLOOR TODAY", 16, BLUE, bold=True)
FLOOR = ["Changing welding consumables",
         "Adjusting the key welding settings within their ranges",
         "Closing ANDON calls"]
for k, t in enumerate(FLOOR):
    y = 2.62 + k * 0.80
    rect(s, 9.25, y + 0.09, 0.14, 0.14, fill=BLUE, shape=MSO_SHAPE.CHEVRON)
    text(s, 9.55, y, 3.3, 0.7, t, 14, NAVY)
text(s, 9.25, 5.45, 3.6, 0.3, "Next: Level I exam, week 40.", 12,
     BLUE, bold=True)
notes(s,
      "11 achieved, 6 in progress, 6 not started, out of 23.",
      "Level I: 5 of 9. Level L: the five skills practiced with support.",
      "U and O are mostly ahead; 4M is my first step in O.")

# ============================================================ 8 · industrialization
# Two different numbers, both kept. Diego's 60 / 40 is how his time on the
# floor is split. The supervisor's 90 % is how ready the MY27 Cummins DOC and
# DPF program is: close to launch, planned for January 2027.
s = S[7]
set_title(s, "Most of my time goes to a program that is 90% ready",
          "My time on the floor, and where the program stands")
retext(s, "TextBox 4", "60%")
retext(s, "TextBox 5", "of my time: industrialization, with Engineering")
shape(s, "Rectangle 7").width = int(shape(s, "Rectangle 6").width * 0.6)
retext(s, "TextBox 8", "40%")
retext(s, "TextBox 12", "Industrialization and successful PPAP runs of the "
       "MY27 Cummins DOC and DPF program, together with Engineering.")
shape(s, "TextBox 12").height = Inches(0.9)
retext(s, "TextBox 14", "PROGRAM PROGRESS")
lbl = shape(s, "TextBox 15")                 # the definition makes way
lbl._element.getparent().remove(lbl._element)
text(s, 6.20, 4.28, 1.9, 0.8, "90%", 40, LIME, bold=True)
tb, tf = txbox(s, 8.05, 4.36, 4.60, 0.7)
para(tf, "ready to launch", 14, WHITE, bold=True, first=True)
para(tf, "Start of production planned for January 2027.", 12, GRAY,
     space_before=2)
rect(s, 6.20, 5.30, 6.45, 0.16, fill=GRAY, opacity=30)
rect(s, 6.20, 5.30, 6.45 * 0.9, 0.16, fill=LIME)
notes(s,
      "Two numbers, not to be mixed up.",
      "60% of my time goes to industrialization, 40% to support at the cells.",
      "The program itself is 90% ready, with launch planned for January.",
      "The cells are where most of the hands-on learning happens.")

# ============================================================ 9 · 4M idea
s = S[8]
retext(s, "TextBox 5", "Long weld seams were hard to control by hand, the "
       "welder worked in uncomfortable positions, and welding alone took "
       "146 s per part.")
shape(s, "TextBox 5").height = Inches(0.95)
for nm, y in (("TextBox 6", 3.50), ("TextBox 7", 3.88)):
    shape(s, nm).top = Inches(y)
# the PHOTO label now sits hidden behind Diego's photo
lbl = shape(s, "TextBox 13")
lbl._element.getparent().remove(lbl._element)

# ============================================================ 12 · next
s = S[11]
retext(s, "TextBox 7", "October, validation on the floor")
retext(s, "TextBox 10", "Create and set up robot programs on my own "
       "(SKS and Miller).")
retext(s, "TextBox 11", "November to December")
retext(s, "TextBox 15", "November to December, already started with 4M")
retext(s, "TextBox 19", "Level I exam confirmed, week 40")
retext(s, "TextBox 21", "Supervised time programming robots in SKS and "
       "Miller")
notes(s,
      "L in October, U and O in November and December.",
      "Ask clearly: the exam in week 40, supervised programming time, "
      "feedback on the proposal.")

# ============================================================ 13 · conclusions
s = S[12]
retext(s, "TextBox 7", "I studied all of Level I and mastered five of its "
       "nine topics, and I practiced the five Level L skills with my tutor "
       "and the technicians.")
shape(s, "TextBox 7").height = Inches(0.85)
retext(s, "TextBox 10", "Close the four open Level I topics and pass the "
       "exam in week 40.")
retext(s, "TextBox 11", "Validate Level L on the floor in October, then "
       "reach U and O before January.")
for nm in ("TextBox 10", "TextBox 11"):
    shape(s, nm).height = Inches(0.6)
notes(s,
      "In place: 5 of 9 Level I topics, Level L practiced, 4M running, "
      "training proposal.",
      "Next: exam in week 40, Level L in October, U and O by January.",
      "Close with the last line. Pause, then questions.")

prs.save(OUT)
print("saved", OUT)
