# -*- coding: utf-8 -*-
"""Cronograma ILUO de soldadura, en el formato "Plan de actividades" de planta.

Conserva la gramática de la plantilla original: encabezado azul 0070C0, barras
verdes 00B050, amarillo para lo que se hace fuera de línea, y un par de filas
Plan / Real por actividad. Cambia la granularidad de día a semana porque el
periodo es de 22 semanas y no de 6.
"""
import datetime as dt
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Cronograma_ILUO_Soldadura_Diego_de_Leon.xlsx")

# ---------------------------------------------------------------- paleta
# tomada tal cual de Cronograma_capacitacion_lideres.xlsx
AZUL      = "FF0070C0"   # encabezados
VERDE     = "FF00B050"   # barras
AMARILLO  = "FFFFFF00"   # fuera de línea
AZUL_CL   = "FF9DC3E6"   # accent1 al 40 %, "en línea"
GRIS      = "FFF2F2F2"
BLANCO    = "FFFFFFFF"
FUENTE    = "Calibri"

fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
thin  = Side(style="thin",   color="FFBFBFBF")
med   = Side(style="medium", color="FF404040")
BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------------------------------------------------------------- calendario
INICIO = dt.date(2026, 8, 3)     # lunes de la semana en que arrancó
FIN    = dt.date(2027, 1, 3)
HOY    = dt.date(2026, 9, 22)
semanas = []
d = INICIO
while d < FIN:
    semanas.append(d)
    d += dt.timedelta(days=7)
NSEM = len(semanas)
MES = {8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre",
       12: "Diciembre", 1: "Enero"}
IDX_HOY = max(i for i, s in enumerate(semanas) if s <= HOY)

# ---------------------------------------------------------------- contenido
# (nivel, actividad, estado, plan[(a,b)...], real[(a,b)...], fuera_de_linea)
# los índices de semana son 0 = 3-ago-2026
P = lambda *r: list(r)
ACT = [
 ("BANNER", "NIVEL I — FUNDAMENTOS DE SOLDADURA  ·  9 habilidades  ·  4 h  ·  examen teórico 8.0"),
 ("I", "Seguridad de soldadura y de celda robótica", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Equipo GMAW: fuente, alimentador, antorcha, consumibles", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Mecanismos de transferencia de metal", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Variables del proceso y variables esenciales (CES §9)", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Discontinuidades de GMAW y límites por clase (CES §16)", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Geometría de las juntas soldadas", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Posiciones de soldadura", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Simbología de soldadura (AWS A2.4 e ISO 2553)", "Parcial",
  P((8,10)), P((0,7)), True),
 ("I", "Documentos de control: WPS, PQR, hoja de parámetros", "Parcial",
  P((8,10)), P((0,7)), True),
 ("HITO", "Examen teórico, mínimo 8.0", "Pendiente", P((11,11)), P(), True),

 ("BANNER", "NIVEL L — EJECUCIÓN EN LA CELDA  ·  5 habilidades  ·  8 h  ·  validación en piso"),
 ("L", "Validación de parámetros contra WPS y hoja de parámetros", "Parcial",
  P((9,16)), P((3,7)), False),
 ("L", "Identificación de discontinuidades y defectos", "Parcial",
  P((9,16)), P((3,7)), False),
 ("L", "Ajuste de variables esenciales: corriente, WFS, voltaje, velocidad", "Dominado",
  P((0,7)), P((1,7)), False),
 ("L", "Interpretación de simbología de soldadura en plano", "Dominado",
  P((0,7)), P((2,7)), False),
 ("L", "Cambio de consumibles: punta, liner, tobera, difusor, rodillos", "Dominado",
  P((0,7)), P((0,7)), False),
 ("HITO", "Validación en piso del nivel L", "Pendiente", P((17,17)), P(), False),

 ("BANNER", "NIVEL U — PROGRAMACIÓN  ·  3 habilidades  ·  8 h  ·  examen práctico"),
 ("U", "Creación de programa nuevo de soldadura", "Parcial",
  P((9,14)), P((5,7)), False),
 ("U", "Crear programa en SKS", "No iniciado", P((12,18)), P(), False),
 ("U", "Asignar parámetros de soldadura (SKS / Miller)", "No iniciado",
  P((12,18)), P(), False),
 ("HITO", "Examen práctico del nivel U", "Pendiente", P((19,19)), P(), False),

 ("BANNER", "NIVEL O — ENSEÑAR Y MEJORAR  ·  6 habilidades  ·  20 h  ·  validación de implementación"),
 ("O", "Capacitar a personal nivel 2 o 3", "No iniciado", P((19,21)), P(), False),
 ("O", "Implementación de ideas de mejora", "No iniciado", P((15,20)), P(), False),
 ("O", "Problem solving: 7 básicos, 5W+2H, 8D, Ishikawa", "No iniciado",
  P((17,20)), P(), True),
 ("O", "GD&T", "No iniciado", P((17,20)), P(), True),
 ("O", "Cierre de órdenes del ANDON", "Dominado", P((0,21)), P((1,7)), False),
 ("O", "Lean Manufacturing", "No iniciado", P((17,20)), P(), True),
 ("HITO", "Validación de implementación del nivel O", "Pendiente",
  P((21,21)), P(), False),

 ("BANNER", "FUERA DE LA MATRIZ ILUO"),
 ("X", "Proyecto de integración con ingeniería", "En curso",
  P((0,16)), P((0,7)), False),
]

# ---------------------------------------------------------------- hoja
wb = Workbook()
ws = wb.active
ws.title = "Plan de actividades"
ws.sheet_view.showGridLines = False

C_ITEM, C_ACT, C_NIV, C_EST, C_RESP, C_ST = 2, 3, 4, 5, 6, 7
C0 = 8                                   # primera columna de semana
ws.column_dimensions[CL(C_ITEM)].width = 4.7
ws.column_dimensions[CL(C_ACT)].width  = 62
ws.column_dimensions[CL(C_NIV)].width  = 6.5
ws.column_dimensions[CL(C_EST)].width  = 12
ws.column_dimensions[CL(C_RESP)].width = 26
ws.column_dimensions[CL(C_ST)].width   = 6.5
for i in range(NSEM):
    ws.column_dimensions[CL(C0 + i)].width = 4.3

def put(r, c, v=None, bg=None, bold=False, size=11, color=None,
        align="left", wrap=False, border=True):
    cell = ws.cell(r, c)
    if v is not None:
        cell.value = v
    cell.font = Font(name=FUENTE, size=size, bold=bold,
                     color=color or "FF000000")
    if bg:
        cell.fill = fill(bg)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               wrap_text=wrap)
    if border:
        cell.border = BORDE
    return cell

ULT = C0 + NSEM - 1

# --- título
ws.row_dimensions[2].height = 24
ws.merge_cells(start_row=2, start_column=C_ITEM, end_row=2, end_column=C_ST)
put(2, C_ITEM, "Nombre del proyecto: Certificación ILUO en soldadura robótica GMAW",
    bold=True, size=12, border=False)
ws.merge_cells(start_row=2, start_column=C0, end_row=2, end_column=ULT)
put(2, C0, "Diego Adair de León Márquez  ·  Tenneco Aguascalientes  ·  "
           "1 ago 2026 a 1 ene 2027", bg=AZUL_CL, bold=True, align="center")

# --- encabezados
FIL_MES, FIL_SEM, FIL0 = 4, 5, 6
ws.row_dimensions[FIL_MES].height = 16
ws.row_dimensions[FIL_SEM].height = 26
for c, txt in ((C_ITEM, "Item"), (C_ACT, "Actividad"), (C_NIV, "Nivel"),
               (C_EST, "Estado"), (C_RESP, "Responsable"), (C_ST, "Status")):
    ws.merge_cells(start_row=FIL_MES, start_column=c, end_row=FIL_SEM, end_column=c)
    put(FIL_MES, c, txt, bg=AZUL, bold=True, color=BLANCO, align="center", wrap=True)
    put(FIL_SEM, c, None, bg=AZUL, border=True)

# banda de meses
i = 0
while i < NSEM:
    m = semanas[i].month
    j = i
    while j + 1 < NSEM and semanas[j + 1].month == m:
        j += 1
    ws.merge_cells(start_row=FIL_MES, start_column=C0 + i,
                   end_row=FIL_MES, end_column=C0 + j)
    put(FIL_MES, C0 + i, MES[m], bg=AZUL, bold=True, color=BLANCO,
        align="center", size=10)
    for k in range(i, j + 1):
        put(FIL_MES, C0 + k, None, bg=AZUL)
    i = j + 1
# banda de semanas
for i, s in enumerate(semanas):
    put(FIL_SEM, C0 + i, s.strftime("%d\n%b").replace("Aug", "ago")
        .replace("Sep", "sep").replace("Oct", "oct").replace("Nov", "nov")
        .replace("Dec", "dic").replace("Jan", "ene"),
        bg=AZUL, bold=True, color=BLANCO, align="center", size=7.5, wrap=True)

# --- filas
NIV_BG = {"I": AZUL_CL, "L": AZUL_CL, "U": AZUL_CL, "O": AZUL_CL,
          "X": GRIS, "HITO": GRIS}
r = FIL0
item = 0
for row in ACT:
    if row[0] == "BANNER":
        ws.merge_cells(start_row=r, start_column=C_ITEM, end_row=r, end_column=ULT)
        ws.row_dimensions[r].height = 18
        put(r, C_ITEM, row[1], bg=AZUL, bold=True, color=BLANCO, size=10)
        for c in range(C_ITEM, ULT + 1):
            put(r, c, None, bg=AZUL)
        r += 1
        continue

    niv, act, est, plan, real, fuera = row
    hito = niv == "HITO"
    item += 1
    for c in range(C_ITEM, C_ST):          # Status queda sin combinar
        ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
    ws.row_dimensions[r].height = 15
    ws.row_dimensions[r + 1].height = 15

    put(r, C_ITEM, item, align="center")
    put(r, C_ACT, act, bg=(AMARILLO if fuera else AZUL_CL) if not hito else GRIS,
        bold=hito, wrap=True)
    put(r, C_NIV, "" if hito else niv, bg=GRIS if hito else None,
        align="center", bold=True)
    put(r, C_EST, est, align="center", size=9,
        bg=GRIS if hito else None)
    put(r, C_RESP, "Fernando Robledo (tutor de planta)" if hito
        else "Diego de León", size=9, bg=GRIS if hito else None)
    put(r, C_ST, "Plan", align="center", size=9, bold=True)
    put(r + 1, C_ST, "Real", align="center", size=9, bold=True)
    for c in (C_ITEM, C_ACT, C_NIV, C_EST, C_RESP):
        put(r + 1, c, None, border=True)

    for rr, tramos, color in ((r, plan, AZUL_CL), (r + 1, real, VERDE)):
        for c in range(C0, ULT + 1):
            put(rr, c, None)
        for a, b in tramos:
            for k in range(a, b + 1):
                put(rr, C0 + k, None, bg=AZUL if (hito and rr == r) else color)
    r += 2

FIN_DATOS = r - 1

# --- marca de HOY
col_hoy = C0 + IDX_HOY
for rr in range(FIL_MES, FIN_DATOS + 1):
    cell = ws.cell(rr, col_hoy)
    b = cell.border
    cell.border = Border(left=med, right=b.right, top=b.top, bottom=b.bottom)
put(3, col_hoy, "HOY", bold=True, size=8, align="center", border=False)

# --- leyenda
r += 1
put(r, C_ACT, "Leyenda", bold=True, border=False)
leyenda = [(AZUL_CL, "Plan  ·  y actividad impartida en línea"),
           (VERDE,   "Real  ·  lo que efectivamente ocurrió"),
           (AMARILLO,"Actividad que se imparte fuera de línea, en aula"),
           (AZUL,    "Hito de evaluación")]
for col, txt in leyenda:
    r += 1
    put(r, C_ITEM, None, bg=col)
    put(r, C_ACT, txt, border=False, size=10)

# --- resumen con fórmulas
r += 2
FIL_RES = r
put(r, C_ACT, "Avance por nivel", bold=True, border=False)
r += 1
for c, h in ((C_ACT, "Nivel"), (C_NIV, "Total"), (C_EST, "Dominado"),
             (C_RESP, "Parcial"), (C_ST, "Sin\niniciar")):
    put(r, c, h, bg=AZUL, bold=True, color=BLANCO, align="center", size=9,
           wrap=True)
niv_col, est_col = CL(C_NIV), CL(C_EST)
rango_n = "%s%d:%s%d" % (niv_col, FIL0, niv_col, FIN_DATOS)
rango_e = "%s%d:%s%d" % (est_col, FIL0, est_col, FIN_DATOS)
for niv, etiqueta in (("I", "Nivel I"), ("L", "Nivel L"),
                      ("U", "Nivel U"), ("O", "Nivel O")):
    r += 1
    put(r, C_ACT, etiqueta, size=10)
    put(r, C_NIV, '=COUNTIF(%s,"%s")' % (rango_n, niv), align="center", size=10)
    for c, est in ((C_EST, "Dominado"), (C_RESP, "Parcial"),
                   (C_ST, "No iniciado")):
        put(r, c, '=COUNTIFS(%s,"%s",%s,"%s")' % (rango_n, niv, rango_e, est),
            align="center", size=10)
r += 1
put(r, C_ACT, "Total de la matriz", bold=True, size=10)
put(r, C_NIV, '=SUM(%s%d:%s%d)' % (niv_col, r - 4, niv_col, r - 1),
    bold=True, align="center", size=10)
for c in (C_EST, C_RESP, C_ST):
    put(r, c, '=SUM(%s%d:%s%d)' % (CL(c), r - 4, CL(c), r - 1),
        bold=True, align="center", size=10)

r += 2
put(r, C_ACT, "Fuente de los estados: autoevaluación de Diego de León al "
              "22 de septiembre de 2026, pendiente de validar con el tutor de planta.",
    size=9, border=False)

ws.freeze_panes = "%s%d" % (CL(C0), FIL0)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = "B2:%s%d" % (CL(ULT), r + 1)
wb.save(OUT)
print("guardado", OUT, "| semanas:", NSEM, "| filas de datos:", FIN_DATOS)
